"""從 data/works.json 產生 Excel（全國、國際各一份）。這是給自己用的檔案，保留作者欄位。
用法：python3 pipeline/build_xlsx.py
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
works = json.load(open(os.path.join(ROOT, 'data', 'works.json'), encoding='utf-8'))
authors = {}
for line in open(os.path.join(ROOT, 'data', 'raw', 'scan.jsonl'), encoding='utf-8'):
    r = json.loads(line)
    if r.get('作者'):
        authors[r['sid']] = '、'.join(x.strip() for x in r['作者'].replace('；', ';').split(';') if x.strip())

STAGE_SRC = {'作品編號': '官方（作品編號）', '官網組別': '官方（組別列表）', '官網科別': '官方（科別欄）', '推測': '推測（依學校名稱）'}
# (欄名, 取值, 欄寬, 是否為本站整理欄位)
COLS = [
    ('sid', lambda w: w['id'], 7, 0),
    ('作品名稱', lambda w: w['title'], 40, 0),
    ('科展類別', lambda w: '全國中小學科展' if w['src'] == '全國' else '臺灣國際科展', 14, 0),
    ('屆次', lambda w: w['edition'], 17, 0),
    ('年份', lambda w: w['year'], 7, 1),
    ('學習階段', lambda w: w.get('stage_detail') or w.get('stage') or '', 9, 1),
    ('學習階段來源', lambda w: STAGE_SRC.get(w.get('stage_src'), ''), 16, 1),
    ('科別', lambda w: w['subject'], 18, 0),
    ('標準領域', lambda w: '、'.join(w['domains']), 14, 1),
    ('得獎情形', lambda w: w['award_raw'], 16, 0),
    ('名次', lambda w: w['rank'] or '', 9, 1),
    ('特別獎', lambda w: '、'.join(w['special']), 14, 1),
    ('得獎資料來源', lambda w: '官方作品資料表' if w.get('award_src') else '科教館網站', 12, 1),
    ('學校名稱', lambda w: w['school'], 28, 0),
    ('縣市', lambda w: w['city'] or '', 8, 1),
    ('指導老師', lambda w: '、'.join(w['teachers']), 16, 0),
    ('作者', lambda w: authors.get(w['id'], ''), 20, 0),
    ('關鍵字', lambda w: '、'.join(w['keywords']), 26, 0),
    ('作品編號', lambda w: w['code'] or '', 9, 1),
    ('類型', lambda w: w['type'], 6, 1),
    ('同作品另一科展 sid', lambda w: '、'.join(map(str, w.get('related', []))), 12, 1),
    ('研究內容節錄', lambda w: '有' if os.path.exists(os.path.join(ROOT, 'site', 'data', 'fulltext', f"{w['id']}.json")) else '', 8, 1),
    ('摘要或動機', lambda w: w['abstract'], 40, 0),
    ('檔案連結', lambda w: w['pdf'] or '', 30, 0),
    ('科教館頁面', lambda w: w['url'], 30, 1),
]

def build(src, name):
    wb = Workbook(); ws = wb.active; ws.title = name[:10]
    ws.append([c[0] for c in COLS])
    for w in works:
        if w['src'] != src:
            continue
        ws.append([ILLEGAL_CHARACTERS_RE.sub('', v) if isinstance(v, str) else v for v in (c[1](w) for c in COLS)])
        for col, key in ((24, 'pdf'), (25, 'url')):
            if w[key]:
                cell = ws.cell(ws.max_row, col); cell.hyperlink = w[key]; cell.font = Font(color='0563C1', underline='single')
    orig, added = PatternFill('solid', fgColor='D9E1F2'), PatternFill('solid', fgColor='E2EFDA')
    for c, (_, _, width, is_added) in zip(ws[1], COLS):
        c.font = Font(bold=True); c.fill = added if is_added else orig; c.alignment = Alignment(horizontal='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
    # 說明頁
    info = wb.create_sheet('欄位說明')
    info.append(['欄位', '說明'])
    for k, v in [
        ('藍色標題', '科教館網站的原始欄位，內容未更動'),
        ('綠色標題', '本站整理的欄位'),
        ('學習階段', '國小、國中、高中、高職；2016 年（第 56 屆）起高中職合併為「高中職」'),
        ('學習階段來源', '官方：取自作品編號或官網組別；推測：依學校名稱與相鄰作品判斷（以官方資料驗證約 99%）'),
        ('標準領域', '依「科別對照表」把歷年不同的科別名稱歸成 12 類'),
        ('名次／特別獎', '由得獎情形拆出；第 56～65 屆以官方作品資料表為準，「未得獎」即官方標示「未列獎項」'),
        ('縣市', '依校名判斷，2010 年改制前的縣名已換成現制；少數校名看不出縣市者留空'),
        ('類型', '「評語」為歷屆評審總評，不是作品'),
        ('研究內容節錄', '「有」表示已從 PDF 抽出研究目的、方法、結果與結論，網站研究包會附上（目前為第 56～65 屆）'),
        ('同作品另一科展 sid', '同一作品也參加了臺灣國際科展（或全國科展）時，另一筆的 sid'),
        ('資料來源', '國立臺灣科學教育館，依政府網站資料開放宣告使用'),
    ]:
        info.append([k, v])
    info.column_dimensions['A'].width = 18; info.column_dimensions['B'].width = 90
    for c in info[1]: c.font = Font(bold=True)
    path = os.path.join(ROOT, f'{name}.xlsx')
    wb.save(path)
    print(path, ws.max_row - 1)

build('全國', '全國中小學科展作品_歷屆')
build('國際', '臺灣國際科展作品_歷屆')
